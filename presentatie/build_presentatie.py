"""Build the SyncJournal live-event deck (python-pptx).

Run from the repo root:  python3 presentatie/build_presentatie.py
Colors come from the light theme in work/syncjournal.html; gold matches the Morani logo.
"""
from datetime import date
from pathlib import Path

from PIL import Image, ImageDraw
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parent.parent
OUT = ROOT / "presentatie"
BUILD = OUT / "build"
BUILD.mkdir(exist_ok=True)

C = {k: RGBColor.from_string(v) for k, v in {
    "bg": "F4F6FA", "surface": "FFFFFF", "ink": "191E2A", "ink2": "525A6B",
    "muted": "6B7385", "line": "E8EBF2", "accent": "3B5BD9", "accent_soft": "EDF0FD",
    "accent_line": "CDD6F7", "gold": "C9A84C", "gold_soft": "F8F1DD", "green": "12946A",
    "green_soft": "E3F4EC", "red": "D64259", "red_soft": "FBE8EC", "black": "0B0B0B",
}.items()}
FONT = "Segoe UI"
FONT_HEAD = "Segoe UI Semibold"

MAANDEN = ["januari", "februari", "maart", "april", "mei", "juni", "juli",
           "augustus", "september", "oktober", "november", "december"]
today = date.today()
DATUM = f"{today.day} {MAANDEN[today.month - 1]} {today.year}"

prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
BLANK = prs.slide_layouts[6]


# ---------- image prep ----------
def rounded_image(src, dst, radius_frac=0.012, border=(232, 235, 242), border_px=3):
    im = Image.open(src).convert("RGBA")
    w, h = im.size
    r = int(min(w, h) * radius_frac * 2.2)
    mask = Image.new("L", (w, h), 0)
    ImageDraw.Draw(mask).rounded_rectangle((0, 0, w - 1, h - 1), r, fill=255)
    im.putalpha(mask)
    ImageDraw.Draw(im).rounded_rectangle((0, 0, w - 1, h - 1), r, outline=border + (255,), width=border_px)
    im.save(dst)
    return dst


def morani_logo(dst):
    # Crop the wordmark out of the 1200x630 OG image and round the corners: a black badge.
    im = Image.open(ROOT / "site/og-image.png").convert("RGBA").crop((330, 215, 870, 425))
    w, h = im.size
    mask = Image.new("L", (w, h), 0)
    ImageDraw.Draw(mask).rounded_rectangle((0, 0, w - 1, h - 1), 26, fill=255)
    im.putalpha(mask)
    im.save(dst)
    return dst


LOGO = morani_logo(BUILD / "morani-badge.png")
DASH = rounded_image(OUT / "screens/dashboard.png", BUILD / "dashboard-rounded.png")


# ---------- primitives ----------
def bg(slide):
    f = slide.background.fill
    f.solid()
    f.fore_color.rgb = C["bg"]


def box(slide, x, y, w, h, fill, line=None, radius=0.08, shape=MSO_SHAPE.ROUNDED_RECTANGLE):
    s = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        s.adjustments[0] = radius
    s.fill.solid()
    s.fill.fore_color.rgb = fill
    if line is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = line
        s.line.width = Pt(1)
    s.shadow.inherit = False
    return s


def text(slide, x, y, w, h, runs, size=24, color="ink", font=FONT, align=PP_ALIGN.LEFT,
         anchor=MSO_ANCHOR.TOP, bold=False):
    """runs: str, or list of paragraphs; a paragraph is str or list of (text, color, font, size)."""
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    for m in ("margin_left", "margin_right", "margin_top", "margin_bottom"):
        setattr(tf, m, 0)
    paras = [runs] if isinstance(runs, str) else runs
    for i, para in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        parts = [(para, color, font, size)] if isinstance(para, str) else para
        for t, col, fnt, sz in parts:
            r = p.add_run()
            r.text = t
            r.font.name = fnt
            r.font.size = Pt(sz)
            r.font.bold = bold
            r.font.color.rgb = C[col]
    return tb


def title(slide, t, kicker=None):
    # The app's page header: a narrow accent bar left of the title ("▍Dashboard").
    box(slide, 0.8, 0.72, 0.09, 0.62, C["accent"], radius=0.5)
    text(slide, 1.08, 0.62, 11.2, 0.8, t, size=36, font=FONT_HEAD, anchor=MSO_ANCHOR.MIDDLE)
    if kicker:
        text(slide, 1.1, 1.42, 11, 0.4, kicker, size=16, color="muted")


def notes(slide, t):
    slide.notes_slide.notes_text_frame.text = t.strip()


def bullets(slide, items, x=0.8, y=2.05, w=11.73, row=0.86, size=26, card=True, marker="accent"):
    if card:
        box(slide, x, y, w, row * len(items) + 0.5, C["surface"], C["line"], radius=0.05)
    for i, it in enumerate(items):
        yy = y + 0.25 + i * row
        box(slide, x + 0.45, yy + row / 2 - 0.08, 0.16, 0.16, C[marker], radius=0.3)
        text(slide, x + 0.85, yy, w - 1.2, row, it, size=size, color="ink", anchor=MSO_ANCHOR.MIDDLE)


def new(t=None, kicker=None):
    s = prs.slides.add_slide(BLANK)
    bg(s)
    if t:
        title(s, t, kicker)
    return s


def section_tag(slide, label, col="accent"):
    soft = {"accent": "accent_soft", "gold": "gold_soft", "green": "green_soft"}[col]
    b = box(slide, 10.9, 0.78, 1.63, 0.5, C[soft], radius=0.5)
    text(slide, 10.9, 0.78, 1.63, 0.5, label, size=14, color="ink2" if col == "gold" else col,
         font=FONT_HEAD, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)


# ---------- slides ----------
# 1. Title
s = new()
box(s, 0, 0, 0.22, 7.5, C["accent"], radius=0, shape=MSO_SHAPE.RECTANGLE)
text(s, 1.2, 2.2, 8, 1.3, [[("Sync", "ink", FONT_HEAD, 72), ("Journal", "accent", FONT_HEAD, 72)]])
text(s, 1.22, 3.6, 8, 0.6, f"Live event · {DATUM}", size=26, color="ink2")
box(s, 1.22, 4.45, 0.9, 0.07, C["gold"], radius=0.5)
text(s, 1.22, 4.75, 8, 0.6, "Denny", size=26, color="ink", font=FONT_HEAD)
s.shapes.add_picture(str(LOGO), Inches(8.7), Inches(2.85), width=Inches(3.6))
text(s, 8.7, 4.35, 3.6, 0.4, "trade-community van Morani", size=14, color="muted", align=PP_ALIGN.CENTER)
notes(s, """
Welkom allemaal! Fijn dat jullie er zijn. Ik ben Denny, de maker van SyncJournal. Vandaag laat ik zien wat de journal is, waarom ik hem gebouwd heb, en ik ga met jullie de diepte in op drie onderwerpen. [aanvullen]
""")

# 2. Agenda
s = new("Welkom, nieuwe leden", "Wat je vandaag ziet")
agenda = [("1", "Wat is SyncJournal, en waarom"), ("2", "Rondgang door de app"),
          ("3", "Diepte: tags, playbook, review"), ("4", "Wat er nog aankomt")]
for i, (n, t) in enumerate(agenda):
    x, y = 0.8 + (i % 2) * 5.98, 2.15 + (i // 2) * 2.25
    box(s, x, y, 5.75, 1.95, C["surface"], C["line"], radius=0.08)
    box(s, x + 0.45, y + 0.6, 0.75, 0.75, C["accent_soft"], radius=0.3)
    text(s, x + 0.45, y + 0.6, 0.75, 0.75, n, size=26, color="accent", font=FONT_HEAD,
         align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    text(s, x + 1.5, y + 0.35, 4.0, 1.25, t, size=24, anchor=MSO_ANCHOR.MIDDLE)
notes(s, """
Voor de nieuwe leden: welkom! Vandaag doe ik vier dingen. Eerst vertel ik wat SyncJournal is en waarom ik het gebouwd heb. Daarna loop ik live met jullie door de app, alle pagina's en de instellingen. Dan ga ik dieper in op drie onderwerpen: tags, de playbook en de review-pagina. En ik sluit af met wat er nog aankomt. Vragen mogen tussendoor, maar aan het eind is er ook tijd. [aanvullen]
""")

# 3. Wat is SyncJournal
s = new("Wat is SyncJournal")
pts = ["Trades automatisch binnen", "Blofin · OKX · Kraken Futures · Hyperliquid", "Of via CSV-import",
       "Analyse van je trades", "Alles lokaal in je browser"]
for i, t in enumerate(pts):
    yy = 2.0 + i * 0.95
    sub = i == 1
    if not sub:
        box(s, 0.95, yy + 0.3, 0.15, 0.15, C["accent"], radius=0.3)
    text(s, 1.35, yy, 4.6, 0.8, t, size=17 if sub else 24,
         color="muted" if sub else "ink", anchor=MSO_ANCHOR.MIDDLE)
pic_w = 6.5
s.shapes.add_picture(str(DASH), Inches(6.05), Inches(1.9), width=Inches(pic_w))
text(s, 6.05, 6.15, pic_w, 0.4, "Dashboard, met voorbeelddata", size=14, color="muted", align=PP_ALIGN.CENTER)
notes(s, """
SyncJournal is een trading-journal: een dagboek voor je trades. Je trades komen automatisch binnen via je exchange. Blofin, OKX, Kraken Futures en Hyperliquid kun je direct koppelen, met een read-only sleutel. Staat jouw exchange er niet tussen, dan lees je je trade-history in als CSV-bestand. Dat werkt voor vrijwel elke exchange.

Daarna rekent de journal voor je uit wat werkt en wat niet: je resultaat, je win-rate, je R, en waar je edge zit.

En belangrijk: alles staat lokaal in je eigen browser. Er gaat niets naar een server van ons. Geen account, geen wachtwoord, en het is gratis. Daarom wel: maak af en toe een back-up. [aanvullen]
""")

# 4. Waarom gebouwd
s = new("Waarom ik het gebouwd heb")
bullets(s, ["Mindset is mijn zwakste punt", "Ik heb structuur en kaders nodig",
            "Traden is een goede spiegel", "De journal geeft die structuur",
            "Voor mezelf, én voor jullie"])
notes(s, """
Zoals de meesten van jullie weten is mindset mijn zwakste punt, en daar werk ik het hardst aan. Ik heb structuur en kaders nodig om mijn ding te kunnen doen. Traden is daarin een goede spiegel; ik heb mezelf er ook beter door leren kennen.

Zo liep ik tegen bepaalde dingen aan (ook privé en op mijn werk) waarvoor ik hulp heb gezocht. Dat begint langzaam zijn vruchten af te werpen.

Om die structuur en kaders vorm te geven heb ik de journal gebouwd. Voor mezelf, maar ook om jullie daarbij te helpen. Inmiddels ben ik echt mega blij met hoe de journal draait, en ik werk er nog wekelijks aan om verbeteringen door te voeren.
""")


# 5-6. Rondgang
def route(slide, rows):
    box(slide, 0.8, 2.0, 11.73, 0.9 * len(rows) + 0.4, C["surface"], C["line"], radius=0.05)
    for i, (n, chips, sub) in enumerate(rows):
        yy = 2.2 + i * 0.9
        text(slide, 1.2, yy, 0.6, 0.9, n, size=18, color="muted", font=FONT_HEAD, anchor=MSO_ANCHOR.MIDDLE)
        xx = 1.8
        for c in chips:
            cw = 0.32 + 0.155 * len(c)
            box(slide, xx, yy + 0.17, cw, 0.56, C["accent_soft"], C["accent_line"], radius=0.5)
            text(slide, xx, yy + 0.17, cw, 0.56, c, size=20, color="accent", font=FONT_HEAD,
                 align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
            xx += cw + 0.18
        text(slide, 7.6, yy, 4.7, 0.9, sub, size=18, color="ink2", anchor=MSO_ANCHOR.MIDDLE)


s = new("Rondgang: de pagina's", "Ik laat het live zien")
section_tag(s, "Live demo")
route(s, [("1", ["Dashboard", "Trades", "Kalender"], "je resultaat"),
          ("2", ["Analytics", "Tendencies"], "je edge (en je lek)"),
          ("3", ["Playbook", "Review"], "leren van je trades"),
          ("4", ["Plan"], "vóór de handelsdag"),
          ("5", ["FAQ", "AI-coach"], "hulp")])
notes(s, """
Spiekbriefje rondgang, ik laat het live zien. [aanvullen]

Dashboard: overzicht van je edge. Je totale balans per account, de kerncijfers (netto P&L, win-rate, profit factor, expectancy, gemiddelde R), je equity-curve, maand-P&L en waar je edge zit. Bovenaan filter je op account, pair, playbook, setup, sessie, datum en meer. Het oog-icoon rechtsboven verbergt je bedragen.

Trades: al je trades op één plek. Alle, gesloten of open/partial. Klik een trade open om hem te bewerken: basis, context (sessie, marktconditie, playbook, setup-grade A/B/C), prijzen en risico, tags, notities en screenshots. Met + Nieuwe trade voer je zelf een trade in.

Kalender: je maand in één oogopslag, welke dagen groen en welke rood waren.

Analytics: de cijfers achter je edge. R-multiples, MFE/MAE, trade-score, SQN, heatmaps.

Tendencies: waar zit je edge, en je lek. Bovenaan je sterkste edge en je grootste lek, daaronder een kruis-tabel (bijvoorbeeld setup × sessie). Cellen met minder dan 8 trades zijn vervaagd.

Playbook: je setups met regels. Per setup rekent de journal win-rate en verwachting uit. Daar ga ik zo dieper op in.

Review: leer van elke trade. Terugkijken met je screenshots, markeren als gereviewd, en delen met je coach. Ook daar kom ik zo op terug.

Plan: dat is TradingPlan, een aparte pagina naast de journal. Je regels, weekplan, dagplan, de poort per trade en de afsluiting na de sessie.

FAQ: veelgestelde vragen, met een zoekbalk.

AI-coach: je persoonlijke trading-assistent. Kom ik aan het eind op terug.
""")

s = new("Rondgang: Instellingen", "Ik laat het live zien")
section_tag(s, "Live demo")
route(s, [("1", ["Accounts"], "koppelingen + CSV-import"),
          ("2", ["Weergave"], "thema, valuta, bedragen"),
          ("3", ["Tags"], "je vaste set tags"),
          ("4", ["AI-coach"], "aan/uit, eigen sleutel"),
          ("5", ["Data", "Updates"], "back-up en versie")])
notes(s, """
Spiekbriefje instellingen. [aanvullen]

Accounts: hier koppel je je exchange (Blofin, OKX, Hyperliquid, Kraken Futures) met een read-only API-key. Hier importeer je ook je trade-history als CSV of XLSX, stel je automatisch synchroniseren en het sync-interval in, en voeg je handmatige accounts toe, bijvoorbeeld voor een prop-firm-challenge.

Weergave: licht of donker thema, bedragen verbergen, valuta, en of je de landingspagina overslaat.

Tags: je vaste set tags per categorie. Daar ga ik zo dieper op in.

AI-coach: de coach aan- of uitzetten, de toon, bedragen verbergen voor de coach, en je eigen API-key.

Data: back-up maken (Exporteer JSON) en terugzetten (Importeer), de Google Drive-backup, voorbeelddata laden of opruimen, veiligheidskopieën, overzetten uit de oude TradeJournal, en de prullenbak met verwijderde trades (30 dagen).

Updates: welke versie je draait, met een link naar de changelog, en de knop Check voor updates.

Er is ook nog een tab Wissen, maar daar blijf je normaal vanaf. Maak altijd eerst een back-up.
""")

# 7-9. Tags
s = new("Tags: wat zijn het?")
section_tag(s, "Tags", "gold")
bullets(s, ["Sleutelwoorden die je aan een trade koppelt", "Ordenen, categoriseren, terugvinden",
            "Consistent taggen = data opbouwen", "Sneller je goede en slechte trades zien"], marker="gold")
notes(s, """
Wat zijn tags? Tags zijn sleutelwoorden die je aan een object koppelt. Het helpt je om data te ordenen, te categoriseren en makkelijker terug te vinden.

Tags koppelen aan trades zorgt ervoor dat je de data van je trades op de juiste manier kunt koppelen. Het consistent koppelen van tags aan trades zorgt ervoor dat je data opbouwt, en dat je sneller jouw goede en slechte trades eruit kunt halen.
""")

s = new("Tags instellen", "Instellingen → Tags")
section_tag(s, "Tags", "gold")
cats = ["Setup type", "Confirmaties", "Timeframe", "Emoties", "Fouten", "Missed-redenen", "Eigen labels"]
box(s, 0.8, 2.05, 11.73, 3.0, C["surface"], C["line"], radius=0.05)
xx, yy = 1.3, 2.5
for c in cats:
    cw = 0.5 + 0.17 * len(c)
    if xx + cw > 12.1:
        xx, yy = 1.3, yy + 1.05
    box(s, xx, yy, cw, 0.7, C["gold_soft"], C["gold"], radius=0.5)
    text(s, xx, yy, cw, 0.7, c, size=22, color="ink", font=FONT_HEAD, align=PP_ALIGN.CENTER,
         anchor=MSO_ANCHOR.MIDDLE)
    xx += cw + 0.25
text(s, 0.8, 5.4, 11.73, 0.6, "Afgestemd op de strategieën van Morani", size=22, color="ink2",
     align=PP_ALIGN.CENTER)
notes(s, """
Je tags kun je instellen via de instellingen-pagina. Daarbij heb ik zoveel mogelijk ingespeeld op de strategieën van Sebas en hoe je deze gestructureerd kunt inzetten.

Je hebt:
- het setup type,
- je confirmaties,
- de timeframe,
- emoties,
- fouten,
- missed-redenen (voor het backtesten),
- en je eigen labels.
""")

s = new("Waarom een vaste set tags?")
section_tag(s, "Tags", "gold")
for i, (hd, col, soft, items) in enumerate([
        ("Andere journals", "red", "red_soft", ["Tags toevoegen in de trade zelf", "Steeds nieuwe tags erbij", "Een wirwar, je doet er niks mee"]),
        ("SyncJournal", "green", "green_soft", ["Eén standaard set tags", "Vooraf goed over nadenken", "Data waar je wél iets mee doet"])]):
    x = 0.8 + i * 5.98
    box(s, x, 2.05, 5.75, 4.4, C["surface"], C["line"], radius=0.06)
    box(s, x, 2.05, 5.75, 0.12, C[col], radius=0, shape=MSO_SHAPE.RECTANGLE)
    text(s, x + 0.45, 2.45, 5, 0.6, hd, size=26, color=col, font=FONT_HEAD)
    for j, it in enumerate(items):
        box(s, x + 0.5, 3.53 + j * 0.9, 0.14, 0.14, C[col], radius=0.3)
        text(s, x + 0.85, 3.25 + j * 0.9, 4.7, 0.7, it, size=21, anchor=MSO_ANCHOR.MIDDLE)
notes(s, """
Bij andere journals kun je in de trade zelf tags toevoegen. Het nadeel dat ik daarin vond, is dat je op dat moment steeds andere tags toevoegt aan de al bestaande lijst van tags. Dat werd een grote wirwar aan tags waar je vervolgens niks mee deed.

Daarom is mijn visie om te werken vanuit een standaard set tags waar jij mee wilt werken. Omdat je dan ook vooraf goed nadenkt over welke tags je wilt toevoegen.
""")

# 10-12. Playbook
s = new("Playbook: waar komt het vandaan?")
section_tag(s, "Playbook")
bullets(s, ["Visie van SMB Capital", "Grote prop firm in New York",
            "Mike Bellafiore: boek “The Playbook”", "Ik keek de YouTube-video's (boek nog niet)"])
notes(s, """
Waarom ik de playbook heb gebouwd. De visie achter de playbook komt van SMB Capital, een grote prop firm in New York. Mike Bellafiore heeft er ook een boek over geschreven: The Playbook. Ik heb het zelf (nog) niet gelezen, maar wel een aantal YouTube-video's erover bekeken.
""")

s = new("Waarom de playbook zo belangrijk is")
section_tag(s, "Playbook")
bullets(s, ["Je best en slechtst presterende strategieën", "Je sterke kanten en je valkuilen",
            "Je edge bewijzen met data", "Uit de emotie, traden op harde criteria",
            "Opschalen? Journal + playbooks"])
notes(s, """
De playbook laat zien wat je best en slechtst presterende strategieën zijn, wat je sterke kanten zijn en waar je valkuilen zitten.

Om je edge te bewijzen heb je data nodig. Zo kom je uit de emotie en ga je traden op harde criteria, onderbouwd door data.

Als je een beetje serieus wilt worden met traden en wilt gaan opschalen, dan heb je naar mijn mening een journal en playbooks nodig.
""")

s = new("Morani's strategieën")
section_tag(s, "Playbook")
bullets(s, ["Gebouwd om Morani's strategieën kwijt te kunnen", "De strategieën werken",
            "Het verschil: hoe jij ze toepast", "Intuïtie, of structuur en kaders",
            "Iedereen is anders"])
notes(s, """
Ik heb de journal zo gebouwd dat je de strategieën van Morani er optimaal in kwijt kunt. De strategieën werken, het verschil zit in hoe jij ze toepast.

De één legt makkelijk de puzzelstukjes vanuit kennis, ervaring en intuïtie. De ander heeft meer structuur en kaders nodig. Iedereen is anders, en ik hoop dat de journal jullie daarin kan ondersteunen.
""")

# 13-14. Review
s = new("Review: terugkijken op je week")
section_tag(s, "Review", "green")
bullets(s, ["Backtesten is belangrijk", "Niet alleen in de chart", "Pak je trades van die week erbij",
            "Hoe ben ik ingestapt?", "Hoe is de trade uitgespeeld?"], marker="green")
notes(s, """
Zoals vaak is benoemd: backtesten is belangrijk. Niet alleen in de chart, maar ook even je trades van die week erbij pakken. Hoe ben ik ingestapt, en hoe zijn ze uitgespeeld?
""")

s = new("De review-pagina")
section_tag(s, "Review", "green")
bullets(s, ["Per trade, zelf reviewen", "Per beeld, samen met de coach", "Snel overzicht om samen te doorlopen",
            "Exporteren of presenteren", "Coach geeft zijn comments"], marker="green")
notes(s, """
De review-pagina biedt de mogelijkheid om dit per trade zelf te doen, of per beeld samen met de coach te doen. En je hebt een snel overzicht om samen met de coach te doorlopen.

Je kunt de trades ook exporteren of presenteren, waarbij Sebas zijn comments kan geven.
""")

# 15. Wat er nog aankomt
s = new("Wat er nog aankomt")
cols = [("AI-coach", "accent", ["Chat over je eigen trades", "Je eigen Anthropic-sleutel", "Journal rekent, coach leest"]),
        ("Funded modus", "gold", ["Challenges en funded accounts", "Apart van je eigen geld", "Wel in je strategie-analyse"]),
        ("Dagplan", "green", ["Meer hierover live", "", ""])]
for i, (hd, col, items) in enumerate(cols):
    x = 0.8 + i * 3.99
    box(s, x, 2.05, 3.75, 4.4, C["surface"], C["line"], radius=0.06)
    box(s, x, 2.05, 3.75, 0.12, C[col], radius=0, shape=MSO_SHAPE.RECTANGLE)
    text(s, x + 0.4, 2.45, 3.1, 0.6, hd, size=26, color="ink", font=FONT_HEAD)
    for j, it in enumerate([t for t in items if t]):
        box(s, x + 0.4, 3.62 + j * 0.95, 0.14, 0.14, C[col], radius=0.3)
        text(s, x + 0.7, 3.24 + j * 0.95, 2.85, 0.9, it, size=20, anchor=MSO_ANCHOR.MIDDLE)
notes(s, """
Tot slot: wat er nog aankomt.

AI-coach: een echte coach waarmee je chat over je eigen trades. Dat werkt met je eigen Anthropic-sleutel. De journal rekent, en de coach leest. Dus de cijfers komen uit de journal zelf, de coach kijkt ernaar en praat er met je over.

Funded modus: prop-firm-challenges en funded accounts houd je apart van je eigen geld, maar ze tellen wel mee in je strategie-analyse.

Dagplan: [aanvullen, Denny vertelt dit zelf]
""")

# 16. Vragen
s = new()
box(s, 0, 0, 0.22, 7.5, C["accent"], radius=0, shape=MSO_SHAPE.RECTANGLE)
text(s, 1.2, 1.0, 10, 1.2, "Vragen?", size=60, font=FONT_HEAD)
box(s, 1.22, 2.3, 0.9, 0.07, C["gold"], radius=0.5)
links = [("X", "x.com/moranitrading"), ("YouTube", "youtube.com/@moranitraden"),
         ("Website", "moranitraden.nl"), ("Journal", "syncjournal.nl")]
for i, (lb, url) in enumerate(links):
    yy = 2.9 + i * 0.85
    text(s, 1.22, yy, 1.8, 0.6, lb, size=20, color="muted", anchor=MSO_ANCHOR.MIDDLE)
    text(s, 3.1, yy, 6.5, 0.6, url, size=26, color="accent" if lb == "Journal" else "ink",
         font=FONT_HEAD, anchor=MSO_ANCHOR.MIDDLE)
s.shapes.add_picture(str(LOGO), Inches(10.4), Inches(5.95), width=Inches(2.2))
notes(s, """
Dat was het van mijn kant. Zijn er vragen? Stel ze gerust. Wil je meer van Morani zien: volg Morani op X en YouTube, en kijk op moranitraden.nl. De journal vind je op syncjournal.nl. Dank jullie wel! [aanvullen]
""")

dst = OUT / "syncjournal-live-event.pptx"
prs.save(dst)
print("saved", dst, len(prs.slides), "slides")
